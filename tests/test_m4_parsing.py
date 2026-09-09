from __future__ import annotations

import io
import logging
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from PIL import Image
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from ntulearn_skill.core import AttachmentId, ContentId, CourseId
from ntulearn_skill.core.models import Coverage
from ntulearn_skill.extractors import (
    ClassificationInput,
    ClassificationRepository,
    ClassificationService,
    RuleBasedMaterialClassifier,
    SemanticType,
)
from ntulearn_skill.parsers import (
    ChunkKind,
    DocxParser,
    FallbackOutput,
    ParsedChunk,
    ParsedPayload,
    ParseOperationError,
    ParserDescriptor,
    ParserLimits,
    ParserOptions,
    ParserRegistry,
    ParseService,
    ParseStatus,
    PdfParser,
    RepresentationKind,
)
from ntulearn_skill.parsers.repository import ParseRepository
from ntulearn_skill.storage import (
    Database,
    DomainRepository,
    ResourceRepository,
    ResourceStore,
    RuntimePaths,
)


@dataclass(frozen=True)
class _Harness:
    paths: RuntimePaths
    database: Database
    resources: ResourceRepository
    store: ResourceStore
    content_id: ContentId


def _harness(tmp_path: Path) -> _Harness:
    paths = RuntimePaths(tmp_path / "private-synthetic-runtime")
    database = Database(paths.database)
    resources = ResourceRepository(database)
    store = ResourceStore(paths, resources)
    assert store.initialize() == 5
    domain = DomainRepository(database)
    course_id = CourseId("synthetic", "m4-course")
    content_id = ContentId("synthetic", "m4-content")
    domain.put_course(course_id, code="PH0000", title="Invented Physics Course")
    domain.put_content_node(
        content_id,
        course_id=course_id,
        handler_kind="document",
        title="Invented materials",
        position=0,
    )
    return _Harness(paths, database, resources, store, content_id)


def _run(database: Database) -> int:
    timestamp = datetime(2027, 1, 1, tzinfo=UTC).isoformat()
    with database.transaction() as connection:
        cursor = connection.execute(
            """INSERT INTO sync_run(mode, requested_scope_json, started_at, status)
            VALUES ('synthetic', '{}', ?, 'RUNNING')""",
            (timestamp,),
        )
        assert cursor.lastrowid is not None
        return int(cursor.lastrowid)


def _pdf(text: str) -> bytes:
    output = io.BytesIO()
    document = canvas.Canvas(output, pagesize=(612, 792), invariant=1)
    document.drawString(72, 720, text)
    document.save()
    return output.getvalue()


def _png() -> io.BytesIO:
    output = io.BytesIO()
    Image.new("RGB", (12, 12), color=(20, 40, 60)).save(output, format="PNG")
    output.seek(0)
    return output


def _pdf_with_form_image() -> bytes:
    output = io.BytesIO()
    document = canvas.Canvas(output, pagesize=(612, 792), invariant=1)
    document.beginForm("InventedVisual")
    document.drawImage(ImageReader(_png()), 0, 0, width=120, height=120)
    document.endForm()
    document.doForm("InventedVisual")
    document.save()
    return output.getvalue()


def _ingest(harness: _Harness, payload: bytes, key: str, filename: str):
    return harness.store.ingest(
        io.BytesIO(payload),
        AttachmentId("synthetic", key),
        content_id=harness.content_id,
        sync_run_key=_run(harness.database),
        display_title=f"Invented {key}",
        original_filename=filename,
    )


def _service(harness: _Harness) -> ParseService:
    return ParseService(
        harness.paths,
        harness.resources,
        ParseRepository(harness.database),
        ParserRegistry((PdfParser(), DocxParser())),
    )


def test_docx_preserves_multiline_paragraph_table_order_and_grid_span(tmp_path: Path) -> None:
    harness = _harness(tmp_path)
    document = Document()
    document.add_paragraph("Invented first line\nInvented second line")
    table = document.add_table(rows=1, cols=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "Invented merged cell"
    document.add_paragraph("Invented final paragraph")
    output = io.BytesIO()
    document.save(output)
    written = _ingest(harness, output.getvalue(), "ordered-docx", "ordered.docx")

    parsed = _service(harness).parse_version(written.version.key)

    assert parsed.document.status is ParseStatus.COMPLETE
    assert [chunk.kind for chunk in parsed.chunks] == [
        ChunkKind.PARAGRAPH,
        ChunkKind.TABLE,
        ChunkKind.PARAGRAPH,
    ]
    assert parsed.chunks[0].native_text == "Invented first line\nInvented second line"
    assert [chunk.locator["element_index"] for chunk in parsed.chunks] == [0, 1, 2]
    table_data = parsed.chunks[1].structured_elements[0]
    assert table_data["rows"][0]["cells"][0]["grid_span"] == 2
    assert all("page" not in " ".join(chunk.locator) for chunk in parsed.chunks)


def test_table_only_docx_remains_searchable_native_text(tmp_path: Path) -> None:
    harness = _harness(tmp_path)
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Invented parameter"
    table.cell(0, 1).text = "Invented value"
    table.cell(1, 0).text = "mobility"
    table.cell(1, 1).text = "120"
    output = io.BytesIO()
    document.save(output)
    written = _ingest(harness, output.getvalue(), "table-docx", "table.docx")

    parsed = _service(harness).parse_version(written.version.key)

    assert len(parsed.chunks) == 1
    assert parsed.chunks[0].kind is ChunkKind.TABLE
    assert "mobility\t120" in parsed.chunks[0].native_text


def test_parse_uses_requested_historical_version_and_bounds_input(tmp_path: Path) -> None:
    harness = _harness(tmp_path)
    first = _ingest(
        harness,
        _pdf("Invented historical text with sufficient native content"),
        "versioned-pdf",
        "versioned.pdf",
    )
    second = _ingest(
        harness,
        _pdf("Invented current text with different sufficient native content"),
        "versioned-pdf",
        "versioned.pdf",
    )
    service = _service(harness)

    historical = service.parse_version(first.version.key)
    bounded = service.parse_version(
        second.version.key,
        options=ParserOptions(
            limits=ParserLimits(maximum_input_bytes=16),
        ),
    )

    assert "historical" in historical.chunks[0].native_text
    assert "current" not in historical.chunks[0].native_text
    assert bounded.document.status is ParseStatus.FAILED
    assert bounded.document.error_code == "parser_input_rejected"


def test_classification_is_format_independent_auditable_and_input_versioned(
    tmp_path: Path,
) -> None:
    harness = _harness(tmp_path)
    pdf = _ingest(harness, _pdf("Invented body"), "classify-pdf", "opaque.pdf")
    document = Document()
    document.add_paragraph("Invented body")
    output = io.BytesIO()
    document.save(output)
    docx = _ingest(harness, output.getvalue(), "classify-docx", "opaque.docx")
    repository = ClassificationRepository(harness.database)
    service = ClassificationService(repository, RuleBasedMaterialClassifier())

    pdf_input = ClassificationInput(
        pdf.resource.key,
        pdf.version.key,
        "Week 3 Lecture Slides",
        "opaque-file",
    )
    docx_input = ClassificationInput(
        docx.resource.key,
        docx.version.key,
        "Week 3 Lecture Slides",
        "opaque-file",
    )
    pdf_result = service.classify(pdf_input)
    repeated = service.classify(pdf_input)
    docx_result = service.classify(docx_input)
    revised = service.classify(
        ClassificationInput(
            pdf.resource.key,
            pdf.version.key,
            "Assignment Brief",
            "opaque-file",
        )
    )

    assert pdf_result.run.candidates[0].semantic_type is SemanticType.LECTURE_SLIDES
    assert docx_result.run.candidates[0].semantic_type is SemanticType.LECTURE_SLIDES
    assert repeated.cache_hit
    assert revised.run.key != pdf_result.run.key
    assert revised.run.input_hash != pdf_result.run.input_hash
    assert revised.run.candidates[0].semantic_type is SemanticType.ASSIGNMENT_BRIEF
    assert pdf_result.run.candidates[0].evidence[0].reference == "title"
    assert "Week 3" not in repr(pdf_result.run.candidates[0].evidence)


def test_unknown_binary_records_terminal_unsupported_state(tmp_path: Path) -> None:
    harness = _harness(tmp_path)
    written = _ingest(harness, b"synthetic opaque bytes", "opaque", "opaque.bin")

    parsed = _service(harness).parse_version(written.version.key)

    assert parsed.document.status is ParseStatus.UNSUPPORTED
    assert parsed.document.error_code == "format_unsupported"
    assert parsed.chunks == ()


def test_parser_rejects_linked_blob_shard_without_caching_failure(tmp_path: Path) -> None:
    harness = _harness(tmp_path)
    written = _ingest(
        harness,
        _pdf("Invented symlink boundary text"),
        "linked-shard",
        "linked.pdf",
    )
    blob = harness.paths.root / written.version.blob_relpath
    outside = tmp_path / "outside-shard"
    blob.parent.rename(outside)
    blob.parent.symlink_to(outside, target_is_directory=True)

    service = _service(harness)
    try:
        service.parse_version(written.version.key)
    except ParseOperationError as error:
        assert str(error) == "private parser path contains a symbolic link"
    else:
        raise AssertionError("linked blob shard should be rejected")

    with harness.database.connect() as connection:
        count = connection.execute("SELECT COUNT(*) FROM parsed_document").fetchone()[0]
    assert count == 0
    assert (outside / written.version.sha256).is_file()


def test_picture_only_docx_is_partial_and_requests_fallback(tmp_path: Path) -> None:
    harness = _harness(tmp_path)
    document = Document()
    document.add_picture(_png())
    output = io.BytesIO()
    document.save(output)
    written = _ingest(harness, output.getvalue(), "picture-docx", "picture.docx")

    parsed = _service(harness).parse_version(written.version.key)

    assert parsed.document.status is ParseStatus.PARTIAL
    assert parsed.document.warning_codes == ("docx_visual_evidence_requires_fallback",)
    assert parsed.chunks[0].native_text == ""
    assert parsed.chunks[0].diagnostic is not None
    assert parsed.chunks[0].diagnostic.image_count == 1
    assert parsed.chunks[0].diagnostic.fallback_recommended


def test_pdf_finds_image_nested_inside_form_xobject(tmp_path: Path) -> None:
    harness = _harness(tmp_path)
    written = _ingest(harness, _pdf_with_form_image(), "form-image", "form.pdf")

    parsed = _service(harness).parse_version(written.version.key)

    assert parsed.document.status is ParseStatus.PARTIAL
    diagnostic = parsed.chunks[0].diagnostic
    assert diagnostic is not None
    assert diagnostic.form_xobject_count >= 1
    assert diagnostic.image_count >= 1
    assert diagnostic.fallback_recommended


class _NoisyFallback:
    def __init__(self, canary: str) -> None:
        self.canary = canary

    def represent(self, request):
        print(self.canary)
        logging.getLogger("synthetic.fallback").error("%s %s", self.canary, request.source_path)
        return FallbackOutput(
            RepresentationKind.OCR_TEXT,
            "Invented recovered content",
            None,
            "synthetic-ocr",
            None,
            "synthetic-1",
            "e" * 64,
            0.8,
        )


def test_fallback_stdout_stderr_and_existing_log_handler_are_suppressed(
    tmp_path: Path, capsys
) -> None:
    harness = _harness(tmp_path)
    written = _ingest(harness, _pdf_with_form_image(), "noisy-fallback", "noisy.pdf")
    service = _service(harness)
    parsed = service.parse_version(written.version.key)
    canary = "PRIVATE-FALLBACK-CANARY"
    captured_log = io.StringIO()
    handler = logging.StreamHandler(captured_log)
    logging.getLogger("synthetic.fallback").addHandler(handler)
    try:
        representations = service.apply_selective_fallback(
            parsed.document.key, _NoisyFallback(canary)
        )
    finally:
        logging.getLogger("synthetic.fallback").removeHandler(handler)

    captured = capsys.readouterr()
    assert len(representations) == 1
    assert canary not in captured.out + captured.err + captured_log.getvalue()


def test_classifier_does_not_match_keyword_inside_larger_word() -> None:
    result = RuleBasedMaterialClassifier().classify(
        ClassificationInput(1, None, "Multithreading and overhead", "notes.bin")
    )

    assert [candidate.semantic_type for candidate in result.candidates] == [SemanticType.UNKNOWN]


class _FlakyLoggingParser:
    descriptor = ParserDescriptor("synthetic-flaky", "1", "synthetic-1", frozenset({"pdf"}))

    def __init__(self, canary: str) -> None:
        self.calls = 0
        self.canary = canary

    def parse(self, path, version, options):
        self.calls += 1
        print(self.canary)
        logging.getLogger("synthetic.parser").error("%s %s", self.canary, path)
        if self.calls == 1:
            raise RuntimeError(self.canary)
        return ParsedPayload(
            ParseStatus.COMPLETE,
            Coverage.COMPLETE,
            (
                ParsedChunk(
                    0,
                    ChunkKind.PAGE,
                    "Invented repaired parse",
                    {"format": "pdf", "physical_page_index": 0},
                ),
            ),
        )


def test_failed_parse_is_auditable_but_retried_and_raw_logging_is_suppressed(
    tmp_path: Path, capsys
) -> None:
    harness = _harness(tmp_path)
    written = _ingest(harness, _pdf("Invented flaky source"), "flaky", "flaky.pdf")
    canary = "PRIVATE-PARSER-CANARY"
    parser = _FlakyLoggingParser(canary)
    service = ParseService(
        harness.paths,
        harness.resources,
        ParseRepository(harness.database),
        ParserRegistry((parser,)),
    )
    captured_log = io.StringIO()
    handler = logging.StreamHandler(captured_log)
    logging.getLogger("synthetic.parser").addHandler(handler)
    try:
        failed = service.parse_version(written.version.key)
        repaired = service.parse_version(written.version.key)
        cached = service.parse_version(written.version.key)
    finally:
        logging.getLogger("synthetic.parser").removeHandler(handler)

    captured = capsys.readouterr()
    assert failed.document.status is ParseStatus.FAILED
    assert repaired.document.status is ParseStatus.COMPLETE
    assert not repaired.cache_hit
    assert cached.cache_hit
    assert parser.calls == 2
    assert canary not in captured.out + captured.err + captured_log.getvalue()
    with harness.database.connect() as connection:
        statuses = [
            row[0]
            for row in connection.execute("SELECT status FROM parsed_document ORDER BY parse_key")
        ]
    assert statuses == ["FAILED", "COMPLETE"]


def _all_strings(value):
    if isinstance(value, str):
        return [value]
    if isinstance(value, dict):
        return [text for item in value.values() for text in _all_strings(item)]
    if isinstance(value, (list, tuple)):
        return [text for item in value for text in _all_strings(item)]
    return []


def test_docx_table_structured_text_obeys_document_text_limit(tmp_path: Path) -> None:
    harness = _harness(tmp_path)
    document = Document()
    table = document.add_table(rows=2, cols=2)
    for row in table.rows:
        for cell in row.cells:
            cell.text = "x" * 40
    output = io.BytesIO()
    document.save(output)
    written = _ingest(harness, output.getvalue(), "bounded-table", "bounded.docx")

    parsed = _service(harness).parse_version(
        written.version.key,
        options=ParserOptions(limits=ParserLimits(maximum_total_characters=30)),
    )

    assert sum(len(chunk.native_text) for chunk in parsed.chunks) <= 30
    structured_text = [
        text
        for chunk in parsed.chunks
        for text in _all_strings(chunk.structured_elements)
        if set(text) <= {"x"}
    ]
    assert sum(len(text) for text in structured_text) <= 30
    assert parsed.document.status is ParseStatus.PARTIAL
