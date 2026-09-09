from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

import pytest
from docx import Document
from PIL import Image
from pypdf import PdfReader, PdfWriter
from pypdf.constants import PageLabelStyle
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from ntulearn_skill.core import AttachmentId, ContentId, CourseId
from ntulearn_skill.core.models import Coverage
from ntulearn_skill.parsers.docx import DocxParser
from ntulearn_skill.parsers.models import (
    ChunkKind,
    FallbackOutput,
    ParsedChunk,
    ParsedPayload,
    ParserDescriptor,
    ParserOptions,
    ParseStatus,
    RepresentationKind,
)
from ntulearn_skill.parsers.pdf import PdfParser
from ntulearn_skill.parsers.registry import ParserRegistry
from ntulearn_skill.parsers.repository import ParseRepository
from ntulearn_skill.parsers.service import ParseService
from ntulearn_skill.storage import (
    Database,
    DomainRepository,
    ResourceRepository,
    ResourceStore,
    RuntimePaths,
)


@dataclass(frozen=True)
class _Context:
    paths: RuntimePaths
    database: Database
    resources: ResourceRepository
    store: ResourceStore
    content_id: ContentId


@pytest.fixture
def parser_context(tmp_path: Path) -> _Context:
    paths = RuntimePaths(tmp_path / "private-synthetic-runtime")
    database = Database(paths.database)
    domain = DomainRepository(database)
    resources = ResourceRepository(database)
    store = ResourceStore(paths, resources)
    assert store.initialize() >= 4
    course_id = CourseId("synthetic", "parser-adversarial-course")
    content_id = ContentId("synthetic", "parser-adversarial-content")
    domain.put_course(course_id, code="PH0000", title="Example Physics Course")
    domain.put_content_node(
        content_id,
        course_id=course_id,
        handler_kind="document",
        title="Invented parser fixtures",
        position=0,
    )
    return _Context(paths, database, resources, store, content_id)


def _sync_run(database: Database) -> int:
    with database.transaction() as connection:
        cursor = connection.execute(
            """
            INSERT INTO sync_run(mode, requested_scope_json, started_at, status)
            VALUES ('course', '{}', ?, 'RUNNING')
            """,
            (datetime(2027, 1, 1, tzinfo=UTC).isoformat(),),
        )
        assert cursor.lastrowid is not None
        return int(cursor.lastrowid)


def _ingest(
    context: _Context,
    payload: bytes,
    *,
    remote_key: str,
    filename: str,
):
    return context.store.ingest(
        io.BytesIO(payload),
        AttachmentId("synthetic", remote_key),
        content_id=context.content_id,
        sync_run_key=_sync_run(context.database),
        display_title=f"Invented {remote_key}",
        original_filename=filename,
    )


def _pdf_bytes(page_drawers) -> bytes:
    output = io.BytesIO()
    document = canvas.Canvas(output, pagesize=(612, 792), invariant=1)
    for draw in page_drawers:
        draw(document)
        document.showPage()
    document.save()
    return output.getvalue()


def _text_page(text: str):
    def draw(document: canvas.Canvas) -> None:
        document.drawString(72, 720, text)

    return draw


def _image_page(document: canvas.Canvas) -> None:
    pixels = Image.new("RGB", (16, 16), color=(31, 63, 95))
    encoded = io.BytesIO()
    pixels.save(encoded, format="PNG")
    encoded.seek(0)
    document.drawImage(ImageReader(encoded), 72, 600, width=160, height=160)


def _service(context: _Context, *parsers) -> tuple[ParseService, ParseRepository]:
    repository = ParseRepository(context.database)
    service = ParseService(
        context.paths,
        context.resources,
        repository,
        ParserRegistry(tuple(parsers)),
    )
    return service, repository


class _RecordingFallback:
    def __init__(self) -> None:
        self.ordinals: list[int] = []

    def represent(self, request):
        self.ordinals.append(request.chunk.ordinal)
        return FallbackOutput(
            representation_kind=RepresentationKind.OCR_TEXT,
            text="Invented text recovered from the image page",
            artifact_relpath=None,
            method="synthetic-ocr",
            provider=None,
            engine_version="synthetic-1",
            settings_hash="f" * 64,
            confidence=0.9,
        )


def test_image_only_page_is_partial_and_only_it_receives_selective_fallback(
    parser_context: _Context,
) -> None:
    written = _ingest(
        parser_context,
        _pdf_bytes((_text_page("Cover"), _image_page)),
        remote_key="mixed-visual-pdf",
        filename="mixed.pdf",
    )
    service, repository = _service(parser_context, PdfParser())

    result = service.parse_version(written.version.key)

    assert result.document.status is ParseStatus.PARTIAL
    assert result.document.coverage is Coverage.PARTIAL
    assert len(result.chunks) == 2
    title, image = result.chunks
    assert title.diagnostic is not None
    assert "low_native_text" in title.diagnostic.reasons
    assert not title.diagnostic.fallback_recommended
    assert image.native_text == ""
    assert image.diagnostic is not None
    assert image.diagnostic.image_count >= 1
    assert image.diagnostic.fallback_recommended

    fallback = _RecordingFallback()
    representations = service.apply_selective_fallback(result.document.key, fallback)

    assert fallback.ordinals == [1]
    assert len(representations) == 1
    assert representations[0].chunk_key == image.key
    assert repository.get(result.document.key).chunks[1].native_text == ""


def test_pdf_physical_index_and_logical_label_remain_distinct(
    parser_context: _Context,
) -> None:
    source = _pdf_bytes(
        (_text_page("Invented first page body"), _text_page("Invented second page body"))
    )
    reader = PdfReader(io.BytesIO(source))
    writer = PdfWriter()
    writer.append_pages_from_reader(reader)
    writer.set_page_label(
        0,
        1,
        style=PageLabelStyle.DECIMAL,
        prefix="Sheet-",
        start=7,
    )
    labeled = io.BytesIO()
    writer.write(labeled)
    written = _ingest(
        parser_context,
        labeled.getvalue(),
        remote_key="logical-labels",
        filename="labels.pdf",
    )
    service, _ = _service(parser_context, PdfParser())

    result = service.parse_version(written.version.key)

    assert [chunk.locator["physical_page_index"] for chunk in result.chunks] == [0, 1]
    assert [chunk.locator["physical_page_number"] for chunk in result.chunks] == [1, 2]
    assert [chunk.locator["logical_page_label"] for chunk in result.chunks] == [
        "Sheet-7",
        "Sheet-8",
    ]


def test_docx_nested_tables_in_different_cells_have_exact_distinct_paths(
    parser_context: _Context,
) -> None:
    document = Document()
    outer = document.add_table(rows=1, cols=2)
    left = outer.cell(0, 0).add_table(rows=1, cols=1)
    right = outer.cell(0, 1).add_table(rows=1, cols=1)
    left.cell(0, 0).text = "Invented left evidence"
    right.cell(0, 0).text = "Invented right evidence"
    output = io.BytesIO()
    document.save(output)
    written = _ingest(
        parser_context,
        output.getvalue(),
        remote_key="nested-docx-tables",
        filename="nested.docx",
    )
    service, _ = _service(parser_context, DocxParser())

    result = service.parse_version(written.version.key)

    assert len(result.chunks) == 1
    chunk = result.chunks[0]
    assert chunk.kind is ChunkKind.TABLE
    assert chunk.locator == {"element_index": 0, "format": "docx", "table_index": 0}
    assert "page" not in " ".join(chunk.locator)
    outer_data = chunk.structured_elements[0]
    cells = outer_data["rows"][0]["cells"]
    left_table = next(item["table"] for item in cells[0]["contents"] if item["kind"] == "table")
    right_table = next(item["table"] for item in cells[1]["contents"] if item["kind"] == "table")
    assert left_table["table_path"] == [0, 0, 0, 0]
    assert right_table["table_path"] == [0, 0, 1, 0]
    assert left_table["table_path"] != right_table["table_path"]
    assert "Invented left evidence" in chunk.native_text
    assert "Invented right evidence" in chunk.native_text


def test_encrypted_pdf_is_terminally_unsupported(parser_context: _Context) -> None:
    reader = PdfReader(io.BytesIO(_pdf_bytes((_text_page("Invented protected text"),))))
    writer = PdfWriter()
    writer.append_pages_from_reader(reader)
    writer.encrypt("synthetic-password")
    encrypted = io.BytesIO()
    writer.write(encrypted)
    written = _ingest(
        parser_context,
        encrypted.getvalue(),
        remote_key="encrypted-pdf",
        filename="protected.pdf",
    )
    service, _ = _service(parser_context, PdfParser())

    result = service.parse_version(written.version.key)

    assert result.document.status is ParseStatus.UNSUPPORTED
    assert result.document.coverage is Coverage.UNKNOWN
    assert result.document.warning_codes == ("encrypted_pdf_unsupported",)
    assert result.document.error_code is None
    assert result.chunks == ()


class _ExplodingParser:
    descriptor = ParserDescriptor(
        name="synthetic-exploder",
        version="1",
        engine_version="synthetic-1",
        formats=frozenset({"pdf"}),
    )

    def __init__(self, canary: str) -> None:
        self.canary = canary

    def parse(self, path, version, options):
        print(f"stdout {self.canary} {path}")
        raise RuntimeError(f"private parser detail {self.canary} {path}")


def test_parser_exception_and_output_do_not_expose_private_canary(
    parser_context: _Context, capsys: pytest.CaptureFixture[str]
) -> None:
    canary = "PRIVATE-CANARY-DO-NOT-EXPOSE"
    written = _ingest(
        parser_context,
        _pdf_bytes((_text_page("Invented error fixture"),)),
        remote_key="parser-error",
        filename="error.pdf",
    )
    service, _ = _service(parser_context, _ExplodingParser(canary))

    result = service.parse_version(written.version.key)

    captured = capsys.readouterr()
    assert result.document.status is ParseStatus.FAILED
    assert result.document.error_code == "parser_failed"
    assert result.document.warning_codes == ()
    assert result.chunks == ()
    assert canary not in captured.out
    assert canary not in captured.err
    assert str(parser_context.paths.root) not in captured.out
    assert str(parser_context.paths.root) not in captured.err


class _CountingParser:
    def __init__(self, engine_version: str) -> None:
        self.descriptor = ParserDescriptor(
            name="synthetic-counting-parser",
            version="1",
            engine_version=engine_version,
            formats=frozenset({"pdf"}),
        )
        self.calls = 0

    def parse(self, path, version, options):
        self.calls += 1
        return ParsedPayload(
            status=ParseStatus.COMPLETE,
            coverage=Coverage.COMPLETE,
            chunks=(
                ParsedChunk(
                    ordinal=0,
                    kind=ChunkKind.PAGE,
                    native_text="Invented cache identity text",
                    locator={
                        "format": "pdf",
                        "physical_page_index": 0,
                        "logical_page_label": None,
                    },
                ),
            ),
        )


def test_parse_cache_identity_includes_settings_and_engine_version(
    parser_context: _Context,
) -> None:
    written = _ingest(
        parser_context,
        _pdf_bytes((_text_page("Invented cache source"),)),
        remote_key="cache-identity",
        filename="cache.pdf",
    )
    first_parser = _CountingParser("engine-a")
    first_service, repository = _service(parser_context, first_parser)

    first = first_service.parse_version(written.version.key)
    repeated = first_service.parse_version(written.version.key)
    changed_settings = first_service.parse_version(
        written.version.key,
        options=ParserOptions(low_text_character_threshold=49),
    )
    second_parser = _CountingParser("engine-b")
    second_service = ParseService(
        parser_context.paths,
        parser_context.resources,
        repository,
        ParserRegistry((second_parser,)),
    )
    changed_engine = second_service.parse_version(written.version.key)

    assert not first.cache_hit
    assert repeated.cache_hit
    assert repeated.document.key == first.document.key
    assert not changed_settings.cache_hit
    assert changed_settings.document.key != first.document.key
    assert changed_settings.document.settings_hash != first.document.settings_hash
    assert not changed_engine.cache_hit
    assert changed_engine.document.key not in {
        first.document.key,
        changed_settings.document.key,
    }
    assert changed_engine.document.engine_version == "engine-b"
    assert first_parser.calls == 2
    assert second_parser.calls == 1


def test_neighbor_hydration_uses_stored_local_chunks_after_source_removal(
    parser_context: _Context,
) -> None:
    written = _ingest(
        parser_context,
        _pdf_bytes(
            (
                _text_page("Invented page zero has enough native text for stable extraction"),
                _text_page("Invented page one has enough native text for stable extraction"),
                _text_page("Invented page two has enough native text for stable extraction"),
            )
        ),
        remote_key="neighbor-hydration",
        filename="neighbors.pdf",
    )
    service, repository = _service(parser_context, PdfParser())
    parsed = service.parse_version(written.version.key)
    assert parsed.document.status is ParseStatus.COMPLETE

    (parser_context.paths.root / written.version.blob_relpath).unlink()
    (parser_context.paths.root / written.version.browse_relpath).unlink()

    windows = repository.hydrate_matches((parsed.chunks[1].key,), neighbor_count=1)

    assert len(windows) == 1
    assert windows[0].match_chunk_key == parsed.chunks[1].key
    assert [chunk.ordinal for chunk in windows[0].chunks] == [0, 1, 2]
    assert [chunk.locator["physical_page_index"] for chunk in windows[0].chunks] == [0, 1, 2]
